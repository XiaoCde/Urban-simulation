# -*- coding: utf-8 -*-
r"""渲染标定探针：量出 docx 里 1 cm 在 LibreOffice 渲染图上到底是多少像素。

用法：
    <python> tools\probe_docx_scale.py
产出：
    out\报告预览\probe.docx
    out\探针\probe_ruler.png（5 cm 宽的黑色矩形，用来量长度）
"""
import os

from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = r"D:\Users\asus\Desktop\项目"
OUT = os.path.join(ROOT, "out", "报告预览")
RULER = os.path.join(ROOT, "out", "探针", "probe_ruler.png")

CJK60 = "一二三四五六七八九十" * 6
BODY_SIZE = 10.0


def main():
    os.makedirs(os.path.dirname(RULER), exist_ok=True)
    os.makedirs(OUT, exist_ok=True)

    # 5.00 cm 宽、0.50 cm 高的纯黑标尺（150 dpi：295 x 30 px）
    Image.new("RGB", (295, 30), (0, 0, 0)).save(RULER)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.font.size = Pt(BODY_SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    def add(text, size=None, width_cm=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if width_cm is not None:
            p.add_run().add_picture(RULER, width=Cm(width_cm))
        else:
            r = p.add_run(text)
            r.font.name = "微软雅黑"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            r.font.size = Pt(size or BODY_SIZE)
        return p

    add("RULER5CM", width_cm=5.0)
    add("MARK-LEFT")
    add(CJK60)
    add("IIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIIII")

    f = os.path.join(OUT, "probe.docx")
    doc.save(f)
    print("已生成 %s (%d B)" % (f, os.path.getsize(f)))


if __name__ == "__main__":
    main()
